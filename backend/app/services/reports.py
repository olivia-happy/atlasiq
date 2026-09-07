from __future__ import annotations

import os
from pathlib import Path
from uuid import uuid4

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from sqlalchemy import select

from ..data import iso_now
from ..db import ResearchReportRecord, session_scope
from ..schemas import ResearchReport, ScenarioRequest
from .countries import get_country_profile
from .market import get_overview, run_scenario, simulate_uncertainty
from .market_packs import active_pack_version
from .sources import list_public_sources
from .factors import build_factor_snapshot


def reports_directory() -> Path:
    configured = os.getenv("ATLASIQ_REPORT_DIR")
    directory = Path(configured) if configured else Path(__file__).resolve().parents[3] / "generated-reports"
    directory.mkdir(parents=True, exist_ok=True)
    return directory


def build_report_markdown(profile_id: str, scenario: ScenarioRequest) -> tuple[str, float, str, int]:
    from .news import list_news

    profile = get_country_profile(profile_id)
    if profile is None:
        raise ValueError("Country profile not found")
    overview = get_overview(profile_id)
    result = run_scenario(scenario, profile_id)
    simulation = simulate_uncertainty(scenario, profile_id=profile_id)
    factor_snapshot = build_factor_snapshot(profile_id)
    version = active_pack_version(profile.market_pack_id)
    source_lines = [f"- {source.name}: {source.status}; {source.observation_label}" for source in list_public_sources(profile_id)]
    evidence_lines = [f"- {item.id}: {item.title}" for item in overview.evidence]
    news_lines = [f"- {item.title} ({item.source})" for item in list_news(profile_id)[:3]] or ["- No profile-specific live news cached yet."]
    factor_lines = [
        f"- {item.key}: {item.score:.1f}; status={item.status}; coverage={item.coverage_ratio:.0%}; "
        f"sources={', '.join(item.source_ids) or 'none'}; {item.explanation}"
        for item in factor_snapshot.items
    ]
    evidence_ids = sorted({evidence_id for item in factor_snapshot.items for evidence_id in item.evidence_ids})
    stored_evidence_lines = [f"- {evidence_id}" for evidence_id in evidence_ids] or ["- No profile-scoped text evidence is stored in this snapshot."]
    markdown = "\n".join(
        [
            f"# AtlasIQ {profile.name} renewable-energy market brief",
            "",
            "## Scope and audit trail",
            f"- Country profile: {profile.name} (`{profile.id}`)",
            f"- Market Pack: {profile.market_pack_id} version {version}",
            "- Sources: OWID Energy, Open-Meteo, configured RSS/GDELT signals",
            "",
            "## Decision conclusion",
            f"- Baseline score: {overview.score}",
            f"- Recommendation: {overview.recommendation}",
            f"- Scenario score: {result.score} ({result.delta:+.1f} vs baseline)",
            "",
            "## Scenario uncertainty",
            f"- P10 / P50 / P90: {simulation.p10} / {simulation.p50} / {simulation.p90}",
            f"- Probability of reaching research threshold: {simulation.probability_research:.1%}",
            "- This is deterministic Monte Carlo scenario analysis, not an investment forecast.",
            "",
            "## Data freshness",
            *source_lines,
            "",
            "## Data coverage",
            f"- Feature snapshot: {factor_snapshot.id} captured at {factor_snapshot.captured_at}",
            *factor_lines,
            "",
            "## Evidence IDs",
            *stored_evidence_lines,
            "",
            "## Evidence",
            *evidence_lines,
            "",
            "## Market signals requiring human verification",
            *news_lines,
            "",
            "## Disclaimer",
            "This local report supports research discussions only. Validate formal policies, grid access, pricing, and project-site facts before a business decision.",
        ]
    )
    return markdown, overview.score, overview.recommendation, version


def render_docx(markdown: str, target: Path) -> None:
    document = Document()
    for line in markdown.splitlines():
        if line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line:
            document.add_paragraph(line)
    document.save(target)


def render_pdf(markdown: str, target: Path) -> None:
    pdf = canvas.Canvas(str(target), pagesize=A4)
    _, height = A4
    cursor = height - 48
    pdf.setFont("Helvetica", 10)
    for line in markdown.splitlines():
        if cursor < 48:
            pdf.showPage()
            pdf.setFont("Helvetica", 10)
            cursor = height - 48
        if line.startswith("# "):
            pdf.setFont("Helvetica-Bold", 14)
            pdf.drawString(42, cursor, line[2:].encode("latin-1", "replace").decode("latin-1"))
            pdf.setFont("Helvetica", 10)
        elif line.startswith("## "):
            pdf.setFont("Helvetica-Bold", 11)
            pdf.drawString(42, cursor, line[3:].encode("latin-1", "replace").decode("latin-1"))
            pdf.setFont("Helvetica", 10)
        elif line:
            pdf.drawString(42, cursor, line[:110].encode("latin-1", "replace").decode("latin-1"))
        cursor -= 16
    pdf.save()


def to_report(record: ResearchReportRecord) -> ResearchReport:
    return ResearchReport(
        id=record.id,
        profile_id=record.profile_id,
        market_pack_id=record.market_pack_id,
        market_pack_version=record.market_pack_version,
        score=record.score,
        recommendation=record.recommendation,
        markdown=record.markdown,
        created_at=record.created_at,
    )


def create_research_report(profile_id: str, scenario: ScenarioRequest) -> ResearchReport:
    profile = get_country_profile(profile_id)
    if profile is None:
        raise ValueError("Country profile not found")
    report_id = f"RPT-{uuid4().hex[:12]}"
    markdown, score, recommendation, version = build_report_markdown(profile_id, scenario)
    directory = reports_directory()
    docx_path = directory / f"{report_id}.docx"
    pdf_path = directory / f"{report_id}.pdf"
    try:
        render_docx(markdown, docx_path)
        render_pdf(markdown, pdf_path)
        if docx_path.stat().st_size == 0 or pdf_path.stat().st_size == 0:
            raise ValueError("Rendered report file was empty")
    except Exception:
        docx_path.unlink(missing_ok=True)
        pdf_path.unlink(missing_ok=True)
        raise
    record = ResearchReportRecord(
        id=report_id,
        profile_id=profile.id,
        market_pack_id=profile.market_pack_id,
        market_pack_version=version,
        score=score,
        recommendation=recommendation,
        markdown=markdown,
        created_at=iso_now(),
    )
    with session_scope() as session:
        session.add(record)
        session.flush()
        return to_report(record)


def get_research_report(report_id: str) -> ResearchReport | None:
    with session_scope() as session:
        record = session.scalar(select(ResearchReportRecord).where(ResearchReportRecord.id == report_id))
        return to_report(record) if record is not None else None


def read_report_file(report_id: str, format: str) -> Path | None:
    if format not in {"docx", "pdf"} or get_research_report(report_id) is None:
        return None
    candidate = reports_directory() / f"{report_id}.{format}"
    return candidate if candidate.is_file() else None
